import asyncio
import subprocess
import sys
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile

import pytest

from app.services.errors import ServiceError


async def parse(tmp_path, data, filename="sample.txt"):
    from app.ingestion.parsing import parse_document

    path = tmp_path / "source"
    path.write_bytes(data)
    return await parse_document(path, filename)


@pytest.mark.parametrize("extension", ["txt", "md"])
async def test_text_keeps_line_locations_and_offsets(tmp_path, extension):
    result = await parse(
        tmp_path, "\ufeff标题\r\n\r\n型号 A001\r\n版本 001\n".encode(), f"sample.{extension}"
    )
    assert result.text == "标题\n\n型号 A001\n版本 001"
    assert [block.locator for block in result.blocks] == [
        {"kind": "lines", "line_start": 1, "line_end": 1},
        {"kind": "lines", "line_start": 3, "line_end": 4},
    ]
    for block in result.blocks:
        assert result.text[block.start : block.end] == block.text


@pytest.mark.parametrize(
    ("data", "code"),
    [
        (b"", "invalid_document"),
        (b" \n\t", "empty_document"),
        (b"\xff\xfeA\x00", "invalid_document"),
        (b"abc\x00def", "invalid_document"),
        ("a�b".encode(), "invalid_document"),
    ],
)
async def test_text_rejects_empty_non_utf8_and_damaged_content(tmp_path, data, code):
    with pytest.raises(ServiceError) as error:
        await parse(tmp_path, data)
    assert error.value.status == 422
    assert error.value.code == code
    assert "abc" not in str(error.value)


async def test_large_single_paragraph_markdown_preserves_source_offsets(tmp_path):
    text = "公开合成资料 KX-17：版本 003，允许离线使用。" * 10000
    result = await parse(tmp_path, text.encode(), "public-large.md")
    assert result.text == text
    assert len(result.blocks) == 1
    assert result.blocks[0].locator == {"kind": "lines", "line_start": 1, "line_end": 1}
    assert result.blocks[0].start == 0
    assert result.blocks[0].end == len(text)
    assert result.text[result.blocks[0].start : result.blocks[0].end] == text


async def test_markdown_text_budget_is_a_resource_failure(tmp_path):
    with pytest.raises(ServiceError) as error:
        await parse(tmp_path, b"x" * 5_000_001, "public-over-budget.md")
    assert error.value.code == "parse_limit"


def pdf_bytes(*texts, encrypt=False):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    for text in texts:
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
        )
        content = DecodedStreamObject()
        content.set_data(f"BT /F1 12 Tf 30 250 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(content)
    if encrypt:
        writer.encrypt("synthetic-secret")
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


async def test_pdf_uses_actual_pages_and_rejects_scan_encryption_and_page_limit(tmp_path):
    result = await parse(tmp_path, pdf_bytes("A001 version 001", "B002 version 002"), "sample.pdf")
    assert [block.locator for block in result.blocks] == [
        {"kind": "page", "page": 1},
        {"kind": "page", "page": 2},
    ]
    assert "A001" in result.text and "B002" in result.text
    for invalid, code in [
        (pdf_bytes(""), "empty_document"),
        (pdf_bytes("secret", encrypt=True), "unsupported_document"),
        (pdf_bytes(*(["text"] * 101)), "parse_limit"),
        (b"not a pdf", "invalid_document"),
    ]:
        with pytest.raises(ServiceError) as error:
            await parse(tmp_path, invalid, "sample.pdf")
        assert error.value.code == code


@pytest.mark.parametrize("mixed", [False, True])
async def test_pdf_scan_is_explicitly_rejected_instead_of_omitting_the_page(tmp_path, mixed):
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject, NumberObject

    writer = PdfWriter()
    if mixed:
        writer.append(PdfReader(BytesIO(pdf_bytes("PUBLIC SYNTHETIC text page"))))
    page = writer.add_blank_page(width=100, height=100)
    image = DecodedStreamObject()
    image.set_data(bytes([0, 0, 0]))
    image.update(
        {
            NameObject("/Type"): NameObject("/XObject"),
            NameObject("/Subtype"): NameObject("/Image"),
            NameObject("/Width"): NumberObject(1),
            NameObject("/Height"): NumberObject(1),
            NameObject("/ColorSpace"): NameObject("/DeviceRGB"),
            NameObject("/BitsPerComponent"): NumberObject(8),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/XObject"): DictionaryObject(
                {NameObject("/Image1"): writer._add_object(image)}
            )
        }
    )
    content = DecodedStreamObject()
    content.set_data(b"q 100 0 0 100 0 0 cm /Image1 Do Q")
    page[NameObject("/Contents")] = writer._add_object(content)
    output = BytesIO()
    writer.write(output)
    with pytest.raises(ServiceError) as error:
        await parse(tmp_path, output.getvalue(), "public-scan.pdf")
    assert error.value.code == "unsupported_document"


def docx_bytes(complex_table=False):
    from docx import Document

    document = Document()
    document.add_paragraph("合成手册 A001")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "型号"
    table.cell(0, 1).text = "X100"
    table.cell(1, 0).text = "版本"
    table.cell(1, 1).text = "001"
    if complex_table:
        table.cell(0, 0).merge(table.cell(0, 1))
    document.add_paragraph("正文结束")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


async def test_docx_preserves_body_order_and_real_paragraph_table_locations(tmp_path):
    result = await parse(tmp_path, docx_bytes(), "sample.docx")
    assert result.text == "合成手册 A001\n\n型号\tX100\n\n版本\t001\n\n正文结束"
    assert [block.locator for block in result.blocks] == [
        {"kind": "paragraph", "paragraph": 1},
        {"kind": "table", "table": 1, "row": 1},
        {"kind": "table", "table": 1, "row": 2},
        {"kind": "paragraph", "paragraph": 2},
    ]
    for block in result.blocks:
        assert result.text[block.start : block.end] == block.text


@pytest.mark.parametrize("encoding", ["utf-16", "utf-16-le", "utf-16-be"])
async def test_ordinary_docx_with_utf16_xml_stays_supported(tmp_path, encoding):
    output = BytesIO()
    with (
        ZipFile(BytesIO(docx_bytes())) as original,
        ZipFile(output, "w", ZIP_DEFLATED) as archive,
    ):
        for entry in original.infolist():
            data = original.read(entry)
            if entry.filename in {"word/document.xml", "word/_rels/document.xml.rels"}:
                xml = data.decode("utf-8").split("?>", 1)[1]
                data = ('<?xml version="1.0" encoding="utf-16"?>' + xml).encode(encoding)
            archive.writestr(entry, data)
    result = await parse(tmp_path, output.getvalue(), "public-utf16.docx")
    assert result.text == "合成手册 A001\n\n型号\tX100\n\n版本\t001\n\n正文结束"
    assert result.blocks[-1].locator == {"kind": "paragraph", "paragraph": 2}


@pytest.mark.parametrize("structure", ["nested_table", "vertical_merge", "tracked_change"])
async def test_complex_docx_structure_has_explicit_unsupported_code(tmp_path, structure):
    from docx import Document
    from docx.oxml import OxmlElement

    document = Document()
    paragraph = document.add_paragraph("PUBLIC SYNTHETIC complex structure")
    table = document.add_table(rows=2, cols=1)
    if structure == "nested_table":
        table.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "Nested text"
    elif structure == "vertical_merge":
        table.cell(0, 0).merge(table.cell(1, 0))
    else:
        paragraph._p.append(OxmlElement("w:ins"))
    output = BytesIO()
    document.save(output)
    with pytest.raises(ServiceError) as error:
        await parse(tmp_path, output.getvalue(), "public-complex.docx")
    assert error.value.code == "unsupported_document"


async def test_docx_rejects_merged_cells_and_zip_path_traversal(tmp_path):
    with pytest.raises(ServiceError) as merged:
        await parse(tmp_path, docx_bytes(complex_table=True), "sample.docx")
    assert merged.value.code == "unsupported_document"
    output = BytesIO()
    with (
        ZipFile(BytesIO(docx_bytes())) as original,
        ZipFile(output, "w", ZIP_DEFLATED) as archive,
    ):
        for entry in original.infolist():
            archive.writestr(entry, original.read(entry))
        archive.writestr("../outside", "no extraction allowed")
    with pytest.raises(ServiceError) as traversal:
        await parse(tmp_path, output.getvalue(), "sample.docx")
    assert traversal.value.code == "parse_limit"


async def test_docx_rejects_decompression_bomb(tmp_path):
    output = BytesIO()
    with (
        ZipFile(BytesIO(docx_bytes())) as original,
        ZipFile(output, "w", ZIP_DEFLATED) as archive,
    ):
        for entry in original.infolist():
            archive.writestr(entry, original.read(entry))
        archive.writestr("custom/compressed.bin", b"x" * (5 * 1024 * 1024))
    with pytest.raises(ServiceError) as error:
        await parse(tmp_path, output.getvalue(), "sample.docx")
    assert error.value.code == "parse_limit"


@pytest.mark.parametrize("encoding", ["utf-8", "utf-16", "utf-16-le", "utf-16-be"])
@pytest.mark.parametrize("entry_name", ["word/document.xml", "word/_rels/document.xml.rels"])
@pytest.mark.parametrize("entity_declaration", [False, True])
async def test_docx_rejects_dtd_in_utf8_and_utf16_xml(
    tmp_path, encoding, entry_name, entity_declaration
):
    output = BytesIO()
    declared = "utf-8" if encoding == "utf-8" else "utf-16"
    dtd = '<!DOCTYPE public [<!ENTITY marker "PUBLIC SYNTHETIC">]>'
    if not entity_declaration:
        dtd = "<!DOCTYPE public>"
    with (
        ZipFile(BytesIO(docx_bytes())) as original,
        ZipFile(output, "w", ZIP_DEFLATED) as archive,
    ):
        for entry in original.infolist():
            data = original.read(entry)
            if entry.filename == entry_name:
                xml = data.decode("utf-8").split("?>", 1)[1]
                data = (f'<?xml version="1.0" encoding="{declared}"?>' + dtd + xml).encode(
                    encoding
                )
            archive.writestr(entry, data)
    with pytest.raises(ServiceError) as error:
        await parse(tmp_path, output.getvalue(), "public-dtd.docx")
    assert error.value.code == "unsupported_document"


async def test_docx_entry_limit_is_a_resource_failure_not_corruption(tmp_path):
    output = BytesIO()
    with (
        ZipFile(BytesIO(docx_bytes())) as original,
        ZipFile(output, "w", ZIP_DEFLATED) as archive,
    ):
        for entry in original.infolist():
            archive.writestr(entry, original.read(entry))
        for number in range(2001):
            archive.writestr(f"custom/empty-{number}", b"")
    with pytest.raises(ServiceError) as error:
        await parse(tmp_path, output.getvalue(), "public-entry-limit.docx")
    assert error.value.code == "parse_limit"


@pytest.mark.parametrize("damage", ["malformed_xml", "invalid_encoding", "crc_mismatch"])
async def test_damaged_docx_remains_invalid_instead_of_resource_failure(tmp_path, damage):
    output = BytesIO()
    with (
        ZipFile(BytesIO(docx_bytes())) as original,
        ZipFile(output, "w", ZIP_STORED) as archive,
    ):
        for entry in original.infolist():
            data = original.read(entry)
            if entry.filename == "word/document.xml":
                if damage == "malformed_xml":
                    data = b'<document><paragraph>PUBLIC SYNTHETIC</document>'
                elif damage == "invalid_encoding":
                    data = b'<?xml version="1.0" encoding="utf-8"?><document>\xff</document>'
            archive.writestr(entry, data, compress_type=ZIP_STORED)
    data = output.getvalue()
    if damage == "crc_mismatch":
        # Change one stored content byte without updating the archive CRC.
        marker = "合成手册 A001".encode()
        assert data.count(marker) == 1
        index = data.index(marker)
        data = data[:index] + bytes([data[index] ^ 1]) + data[index + 1 :]
    with pytest.raises(ServiceError) as error:
        await parse(tmp_path, data, "public-damaged.docx")
    assert error.value.code == "invalid_document"


async def test_parser_deadline_terminates_worker(tmp_path, monkeypatch):
    from app.ingestion import parsing

    monkeypatch.setattr(parsing, "PARSER_TIMEOUT_SECONDS", 0.001)
    with pytest.raises(ServiceError) as error:
        await parse(tmp_path, b"safe text")
    assert error.value.code == "parse_timeout"


async def test_docx_rejects_nonempty_headers_instead_of_omitting_content(tmp_path):
    from docx import Document

    document = Document()
    document.sections[0].header.paragraphs[0].text = "This is the applicable model A001"
    document.add_paragraph("Only use for the model named in the header")
    output = BytesIO()
    document.save(output)
    with pytest.raises(ServiceError) as error:
        await parse(tmp_path, output.getvalue(), "header.docx")
    assert error.value.code == "unsupported_document"


async def test_pdf_blank_pages_do_not_invent_or_shift_page_numbers(tmp_path):
    result = await parse(tmp_path, pdf_bytes("", "A001"), "pages.pdf")
    assert [block.locator for block in result.blocks] == [{"kind": "page", "page": 2}]


async def test_parse_cancellation_reaps_the_actual_child_process(tmp_path, monkeypatch):
    from app.ingestion import parsing

    real_create = asyncio.create_subprocess_exec
    children = []
    started = asyncio.Event()

    async def capture_process(*args, **kwargs):
        process = await real_create(*args, **kwargs)
        children.append(process)
        started.set()
        return process

    monkeypatch.setattr(parsing.asyncio, "create_subprocess_exec", capture_process)
    task = asyncio.create_task(parse(tmp_path, b"safe text"))
    await started.wait()
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert len(children) == 1 and children[0].returncode is not None


def test_worker_enforces_memory_limit_before_document_processing():
    from app.ingestion import parser_process

    code = (
        "import sys; sys.path.insert(0, sys.argv[1]); import parser_process; "
        "handle = parser_process._limit_resources()\n"
        "try:\n    allocation = bytearray(600 * 1024 * 1024)\n"
        "except MemoryError:\n    sys.exit(23)\n"
        "sys.exit(0)\n"
    )
    result = subprocess.run(
        [sys.executable, "-I", "-c", code, str(Path(parser_process.__file__).parent)],
        capture_output=True,
        timeout=10,
    )
    assert result.returncode == 23
